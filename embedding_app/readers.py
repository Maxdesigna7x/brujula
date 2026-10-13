import os, re, json, subprocess, zipfile
from pathlib import Path
from html import unescape

TEXT={'.txt','.md','.rst','.csv','.tsv','.json','.yaml','.yml','.html','.htm','.pdf','.docx','.odt','.rtf'}
IMAGE={'.jpg','.jpeg','.png','.webp','.bmp','.tif','.tiff','.gif'}
VIDEO={'.mp4','.mkv','.mov','.avi','.webm','.m4v','.mpeg','.mpg'}
AUDIO={'.mp3','.wav','.flac','.m4a','.ogg','.opus','.aac'}
SENSITIVE={'.pem','.key','.p12','.pfx','.kdbx','.env'}
def kind(path):
    e=Path(path).suffix.lower()
    return 'text' if e in TEXT else 'image' if e in IMAGE else 'video' if e in VIDEO else 'audio' if e in AUDIO else 'other'

def select_text(text, words=3000):
    parts=text.split()
    if len(parts)<=words:return ' '.join(parts)
    # Equal coverage across start/middle/end, never silently just truncate the start.
    n=max(1,words//3);mid=len(parts)//2
    return '\n[…]\n'.join([' '.join(parts[:n]),' '.join(parts[max(0,mid-n//2):mid+n//2]),' '.join(parts[-n:])])

def text_samples(path,limit=3000):
    p=Path(path);e=p.suffix.lower()
    if e=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(p);count=len(reader.pages)
        positions=sorted(set([0,max(0,count//2),max(0,count-1)]+list(range(min(3,count)))))
        text='\n'.join((reader.pages[i].extract_text() or '') for i in positions if i<count)
        if not text.strip():raise ValueError('PDF sin texto extraíble. OCR pendiente; se busca por nombre.')
    elif e=='.docx':
        from docx import Document
        doc=Document(p);text='\n'.join(x.text for x in doc.paragraphs)
        text+='\n'+'\n'.join(' '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    elif e=='.odt':
        with zipfile.ZipFile(p) as z:
            if z.getinfo('content.xml').file_size>20_000_000:raise ValueError('Documento ODT demasiado grande.')
            text=re.sub('<[^>]+>',' ',z.read('content.xml').decode('utf8','replace'))
    elif e=='.rtf':
        raise ValueError('RTF todavía no tiene extractor; disponible por nombre.')
    else:
        # Bounded read. Huge text files get byte samples rather than reading all into RAM.
        with p.open('rb') as f:
            size=p.stat().st_size
            if size<=2_000_000:raw=f.read()
            else:
                chunks=[]
                for offset in [0,size//2,max(0,size-600_000)]:f.seek(offset);chunks.append(f.read(600_000))
                raw=b'\n'.join(chunks)
        text=raw.decode('utf8','replace')
        if e in {'.html','.htm'}:
            text=re.sub(r'<(script|style)\b[^>]*>.*?</\1>',' ',text,flags=re.S|re.I)
            text=unescape(re.sub('<[^>]+>',' ',text))
    return select_text(text,limit)

def probe(path):
    out=subprocess.check_output(['ffprobe','-v','error','-show_entries','format=duration:stream=codec_type','-of','json',str(path)],timeout=30)
    info=json.loads(out);duration=float(info.get('format',{}).get('duration',0))
    return duration,any(s.get('codec_type')=='audio' for s in info.get('streams',[]))

def positions(duration,count,window=0):
    if duration<=0:return [0.0]
    end=max(0,duration-window)
    if count<=1 or end==0:return [end/2]
    return [end*(i+.5)/count for i in range(count)]

def extract_media(path,category,settings,directory):
    directory=Path(directory);duration,has_audio=probe(path);items=[]
    if category=='video':
        for i,t in enumerate(positions(duration,int(settings['video_frames']))):
            dest=directory/f'frame-{i}.jpg'
            subprocess.run(['ffmpeg','-nostdin','-v','error','-ss',str(t),'-i',str(path),'-frames:v','1','-vf',"scale='min(640,iw)':-2",'-y',str(dest)],check=True,timeout=60)
            if dest.exists():items.append(('image',str(dest),t))
    if has_audio and settings['enable_audio']:
        count=int(settings['audio_windows']);window=float(settings['audio_seconds'])
        for i,t in enumerate(positions(duration,count,window)):
            dest=directory/f'audio-{i}.wav'
            subprocess.run(['ffmpeg','-nostdin','-v','error','-ss',str(t),'-i',str(path),'-t',str(window),'-vn','-ac','1','-ar','16000','-y',str(dest)],check=True,timeout=60)
            items.append(('audio',str(dest),t))
    return items
